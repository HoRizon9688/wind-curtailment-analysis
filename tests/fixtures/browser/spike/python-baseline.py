"""Print what openpyxl (the Python baseline) sees in each OOXML spike fixture.

Used to compare the browser adapter against `upload_pipeline.table_rows`, which
selects `book['功率预测'] if present else book.active`.
"""
import io
import json
import pathlib
import warnings
import zipfile

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
MAX_UNCOMPRESSED = 150_000_000


def describe(name):
    blob = (HERE / name).read_bytes()
    out = {"file": name, "bytes": len(blob)}
    if not zipfile.is_zipfile(io.BytesIO(blob)):
        out["not_zip"] = True
        return out
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        out["declared_uncompressed"] = sum(i.file_size for i in z.infolist())
        out["has_workbook_xml"] = "xl/workbook.xml" in z.namelist()
    if not out["has_workbook_xml"]:
        out["missing_workbook"] = True
        return out
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        out["sheetnames"] = list(book.sheetnames)
        out["active_title"] = book.active.title
        chosen = book["功率预测"] if "功率预测" in book.sheetnames else book.active
        out["chosen"] = chosen.title
        rows = list(chosen.values)
        out["row_count"] = len(rows)
        out["header"] = [None if v is None else str(v) for v in (rows[0] if rows else [])]
        out["row1"] = [describe_cell(v) for v in (rows[1] if len(rows) > 1 else [])]
    finally:
        book.close()
    return out


def describe_cell(value):
    if isinstance(value, bool):
        return {"type": "bool", "value": value}
    if isinstance(value, (int, float)):
        return {"type": "number", "value": value}
    if hasattr(value, "isoformat"):
        return {"type": "datetime", "value": value.isoformat()}
    if value is None:
        return {"type": "null"}
    return {"type": "string", "value": str(value)}


if __name__ == "__main__":
    targets = [
        "forecast-ooxml.xls",
        "minute-power-ooxml.xlsx",
        "forecast-1904.xlsx",
        "forecast-no-power-sheet.xls",
        "forecast-active-second.xlsx",
        "forecast-active-first.xlsx",
        "bounds-wide.xlsx",
        "no-workbook.xlsx",
        "corrupt.xlsx",
    ]
    for name in targets:
        try:
            print(json.dumps(describe(name), ensure_ascii=False))
        except Exception as exc:  # noqa: BLE001 - report the failure verbatim
            print(json.dumps({"file": name, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
